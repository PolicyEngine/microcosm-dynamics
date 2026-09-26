"""Track M's Census threshold capture, 1982-2022 (plan item M2; cos d279).

Artifact tier: reads the thirty-five Census workbooks committed in
``data/external/census_poverty_thresholds`` (1982, 1986, 1988, 1989, 1991,
1992 and 1994-2022; d194 and d279) and the capture
``data/external/census_poverty_thresholds_1982_2022.json``.  It checks the
pin, re-runs the capture, compares the one-person 65-and-over weighted
average of every year with the workbook cell read without the parser or
openpyxl, and exercises the two layout departures of 2019 and 2022 (empty
extra worksheets, weighted averages rounded to $10) and their refusals.
The departures before 2003, the differential by row label, the cross-check
against Census's HTML Table 1 and the invariants are in
``test_threshold_capture_before_2003.py``.  No PSID file is read.
"""

from __future__ import annotations

import copy
import hashlib
import importlib.util
import json
import re
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pytest

from populace_dynamics.min_benefit_track_m import rules, thresholds

ROOT = Path(__file__).resolve().parents[2]
WORKBOOKS = ROOT / "data" / "external" / "census_poverty_thresholds"
CAPTURE = (
    ROOT / "data" / "external" / "census_poverty_thresholds_1982_2022.json"
)
YEARS = (1982, 1986, 1988, 1989, 1991, 1992, *range(1994, 2023))
_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
)


def _script():
    path = ROOT / "scripts" / "capture_track_u_parameters.py"
    spec = importlib.util.spec_from_file_location("capture_track_m", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def script():
    return _script()


@pytest.fixture(scope="module")
def capture():
    return json.loads(CAPTURE.read_text())


def _name(year: int) -> str:
    return f"thresh{year % 100:02d}.xlsx"


def _first_sheet_cells(path: Path) -> tuple[int, dict[str, object]]:
    """The number of worksheets and the first one's cells, from the XML."""

    with zipfile.ZipFile(path) as archive:
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
        sheets = list(book.iter(f"{_MAIN}sheet"))
        target = targets[sheets[0].attrib[_REL]].lstrip("/")
        target = target if target.startswith("xl/") else f"xl/{target}"
        strings = [
            "".join(t.text or "" for t in item.iter(f"{_MAIN}t"))
            for item in ET.fromstring(
                archive.read("xl/sharedStrings.xml")
            ).iter(f"{_MAIN}si")
        ]
        sheet = ET.fromstring(archive.read(target))
    cells: dict[str, object] = {}
    for cell in sheet.iter(f"{_MAIN}c"):
        assert cell.find(f"{_MAIN}f") is None, "formula cell"
        value = cell.find(f"{_MAIN}v")
        if value is None:
            continue
        if cell.attrib.get("t") == "s":
            cells[cell.attrib["r"]] = strings[int(value.text)]
        else:
            cells[cell.attrib["r"]] = float(value.text)
    return len(sheets), cells


# -------------------------------------------------------------------------
# Pins and reproduction
# -------------------------------------------------------------------------
def test_every_workbook_is_committed_and_pinned(script):
    assert sorted(script.CENSUS_WORKBOOK_SHA256) == sorted(
        _name(year) for year in YEARS
    )
    provenance = (WORKBOOKS / "provenance.md").read_text()
    for name, pin in script.CENSUS_WORKBOOK_SHA256.items():
        digest = hashlib.sha256((WORKBOOKS / name).read_bytes()).hexdigest()
        assert digest == pin, name
        assert f"| `{name}` |" in provenance and f"`{pin}`" in provenance


def test_the_download_record_says_who_fetched_the_workbooks(script):
    """Independent review (2026-09-25): Max approved the downloads (cos
    d194, d279); the orchestrating Claude Code session fetched the files
    with curl and wrote their SHA256SUMS, as its session record shows.
    Earlier text said Max downloaded them.  No record may say so."""

    from populace_dynamics.min_benefit_track_m import policy as pol

    note = pol.MAX_RULINGS["census_threshold_download"]["note"]
    spec_text = (
        ROOT / "docs" / "design" / "minimum_benefits_comparison.md"
    ).read_text()
    texts = {
        "provenance.md": (WORKBOOKS / "provenance.md").read_text(),
        "capture script": Path(script.__file__).read_text(),
        "MAX_RULINGS note": note,
        "M1 specification": spec_text,
    }
    wrong = re.compile(
        r"downloaded\s+by\s+Max|Max\s+downloaded|his\s+`?SHA256SUMS",
        re.IGNORECASE,
    )
    for where, text in texts.items():
        flat = " ".join(text.split())
        assert not wrong.search(flat), where
    assert "orchestrating Claude Code session fetched" in note
    for where in ("provenance.md", "capture script"):
        flat = " ".join(texts[where].split())
        assert "orchestrating Claude Code session" in flat, where


def test_the_capture_is_pinned_and_loads():
    assert thresholds.TRACK_M_THRESHOLDS_PATH == CAPTURE
    assert thresholds.TRACK_M_THRESHOLD_YEARS == YEARS
    digest = hashlib.sha256(CAPTURE.read_bytes()).hexdigest()
    assert digest == thresholds.TRACK_M_THRESHOLDS_SHA256
    loaded = rules.load_aged_thresholds()
    assert sorted(loaded.annual) == list(YEARS)
    assert loaded.source == {
        "kind": "census_capture",
        "path": "data/external/census_poverty_thresholds_1982_2022.json",
        "sha256": thresholds.TRACK_M_THRESHOLDS_SHA256,
        "row": "one_65_plus",
        "years": [1982, 2022],
        "captured_years": list(YEARS),
    }
    # the retired 2003-2022 capture is gone, and nothing reads it
    assert not (
        CAPTURE.parent / "census_poverty_thresholds_2003_2022.json"
    ).exists()
    # rules re-exports the loader and its errors
    assert rules.load_aged_thresholds is thresholds.load_aged_thresholds
    assert rules.AgedThresholds is thresholds.AgedThresholds


def test_the_capture_reproduces_byte_for_byte(script):
    capture = script.build_track_m_threshold_capture(WORKBOOKS)
    text = json.dumps(capture, indent=2) + "\n"
    assert text == CAPTURE.read_text(encoding="utf-8")


def test_the_capture_agrees_with_track_us_where_they_overlap(capture):
    track_u = json.loads(
        (
            ROOT
            / "data"
            / "external"
            / "census_poverty_thresholds_2004_2012.json"
        ).read_text()
    )
    for year in range(2004, 2013):
        key = str(year)
        for block in (
            "weighted_average",
            "weighted_average_all_ages",
            "matrix",
        ):
            assert capture[block][key] == track_u[block][key], (year, block)
        assert capture["sources"][key] == track_u["sources"][key]


def test_capture_records_sources_layouts_and_ratios(capture, script):
    assert capture["years"] == [1982, 2022]
    assert capture["captured_years"] == list(YEARS)
    assert capture["years_not_captured"] == [
        1983,
        1984,
        1985,
        1987,
        1990,
        1993,
    ]
    assert set(capture["decision_records"]) == {
        "d194",
        "d279",
        "d279_earlier_years",
    }
    assert capture["checks"]["workbook_sha256"].startswith("pinned")
    layouts = capture["checks"]["layout_by_year"]
    for year in YEARS:
        source = capture["sources"][str(year)]
        assert source["file"] == _name(year)
        assert source["sha256"] == script.CENSUS_WORKBOOK_SHA256[_name(year)]
        survey = (
            "cps_asec"
            if year >= 2002
            else "cps_ads" if year == 2001 else "march_cps"
        )
        if survey == "cps_asec":
            assert source["cps_asec_year"] == year + 1
            assert "note_survey" not in source
        else:
            assert "cps_asec_year" not in source
            assert source["note_survey"] == survey
            assert source["note_survey_year"] == year + 1
        revision = {
            1982: "Revised on 4/19/2022 due to rounding issues.",
            2000: "Revised on 2/1/2023 due to a formatting error.",
        }.get(year)
        assert source.get("revision_line") == revision, year
        assert ("retrieval" in source) == (year == 1995), year
        expected = {
            "worksheets": 3 if year == 2019 else 1,
            "empty_extra_worksheets": (
                ["Sheet2", "Sheet3"] if year == 2019 else []
            ),
            "weighted_average_unit_dollars": 10 if year == 2022 else 1,
            "note_survey": survey,
            "size_row_noun": "persons" if year == 2001 else "people",
            "revision_line": revision,
        }
        assert layouts[str(year)] == expected, year
    ratios = capture["checks"]["matrix_ratio_by_year_pair"]
    assert list(ratios) == [
        f"{a}-{b}" for a, b in zip(YEARS, YEARS[1:], strict=False)
    ]


@pytest.mark.parametrize("year", YEARS)
def test_the_track_m_row_equals_the_workbook_cell(capture, year):
    """G8's cell: B10, "65 years and over" under one person, read from the
    workbook XML without the parser."""

    n_sheets, cells = _first_sheet_cells(WORKBOOKS / _name(year))
    assert n_sheets == (3 if year == 2019 else 1)
    assert cells["A2"].startswith(f"Poverty Thresholds for {year} by")
    # as printed: a colon after the size-1 label, except 2001's dot
    # leaders; "Weighted Average Thresholds" capitalized in 1988, 1994 and
    # 1998 (compared in lower case)
    assert cells["A8"] == (
        "One person (unrelated individual)......"
        if year == 2001
        else "One person (unrelated individual):"
    )
    assert cells["A10"] == "65 years and over......"
    assert cells["B5"].lower() == "weighted\naverage\nthresholds"
    value = capture["weighted_average"][str(year)]["one_65_plus"]
    assert cells["B10"] == value
    assert rules.load_aged_thresholds().for_year(year) == value
    # 2022 prints the weighted average rounded to $10 beside its single
    # matrix cell; every other year prints the two equal.
    single = capture["matrix"][str(year)]["one_65_plus"]["0"]
    assert cells["C10"] == single
    if year == 2022:
        assert (value, single) == (14040, 14036)
    else:
        assert value == single


# -------------------------------------------------------------------------
# The two layout departures, pinned to their years
# -------------------------------------------------------------------------
def test_2019s_extra_worksheets_are_accepted_only_empty_and_only_in_2019(
    script, tmp_path
):
    path = WORKBOOKS / "thresh19.xlsx"
    with pytest.raises(ValueError, match="3 worksheets"):
        script.read_workbook_rows(path)  # no year: one worksheet only
    sheet, rows = script.read_workbook_rows(path, 2019)
    assert sheet == "Sheet1"
    script.parse_threshold_rows(rows, 2019)
    # a value on an extra worksheet is refused
    copy19 = tmp_path / "thresh19.xlsx"
    shutil.copy2(path, copy19)
    book = openpyxl.load_workbook(copy19)
    book["Sheet3"]["B2"] = 12261
    book.save(copy19)
    with pytest.raises(ValueError, match="'Sheet3' is not empty"):
        script.read_workbook_rows(copy19, 2019)
    # the same worksheets are refused in another year's workbook
    copy18 = tmp_path / "thresh18.xlsx"
    shutil.copy2(WORKBOOKS / "thresh18.xlsx", copy18)
    book = openpyxl.load_workbook(copy18)
    book.create_sheet("Sheet2")
    book.create_sheet("Sheet3")
    book.save(copy18)
    with pytest.raises(ValueError, match="expected 1"):
        script.read_workbook_rows(copy18, 2018)
    assert script.EMPTY_EXTRA_SHEETS == {2019: ("Sheet2", "Sheet3")}


def _rows(script, year: int) -> list[list]:
    return script.read_workbook_rows(WORKBOOKS / _name(year), year)[1]


def _set(rows, row, column, value):
    rows = copy.deepcopy(rows)
    rows[row][column] = value
    return rows


def test_2022s_ten_dollar_weighted_averages(script):
    assert script.WEIGHTED_AVERAGE_UNIT == {2022: 10}
    rows = _rows(script, 2022)
    parsed = script.parse_threshold_rows(rows, 2022)
    assert parsed["weighted_average"]["one_65_plus"] == 14040
    assert parsed["matrix"]["one_65_plus"] == {0: 14036}
    assert all(
        value % 10 == 0
        for value in (
            *parsed["weighted_average"].values(),
            *parsed["all_ages"].values(),
        )
    )
    # 2022 must print every weighted average to $10: B10 at its matrix
    # cell's whole dollars is refused.
    with pytest.raises(ValueError, match=r"not multiples of \$10"):
        script.parse_threshold_rows(_set(rows, 9, 1, 14036), 2022)
    # the $5 slack is half the unit, no more
    with pytest.raises(ValueError, match="outside its matrix cells"):
        script.parse_threshold_rows(_set(rows, 9, 1, 14050), 2022)
    # the same $4 gap in a whole-dollar year is refused
    rows_2021 = _rows(script, 2021)
    assert rows_2021[9][1] == 12996
    with pytest.raises(ValueError, match="outside its matrix cells"):
        script.parse_threshold_rows(_set(rows_2021, 9, 1, 13000), 2021)


def test_every_real_layout_parses_and_moves_by_one_ratio(script, capture):
    previous = None
    for year in YEARS:
        parsed = script.parse_threshold_rows(_rows(script, year), year)
        assert (
            parsed["weighted_average"]
            == capture["weighted_average"][str(year)]
        )
        if previous is not None:
            earlier_year, earlier = previous
            ratio = script.check_matrix_moves_together(
                earlier, parsed["matrix"], year, earlier_year
            )
            key = f"{earlier_year}-{year}"
            assert round(ratio, 6) == pytest.approx(
                capture["checks"]["matrix_ratio_by_year_pair"][key]
            )
        previous = (year, parsed["matrix"])


def test_a_tampered_or_missing_workbook_is_refused(script, tmp_path):
    for path in WORKBOOKS.glob("thresh*.xlsx"):
        shutil.copy2(path, tmp_path / path.name)
    book = openpyxl.load_workbook(tmp_path / "thresh17.xlsx")
    book.worksheets[0]["B10"] = 11757
    book.save(tmp_path / "thresh17.xlsx")
    with pytest.raises(ValueError, match="not the captured Census file"):
        script.build_track_m_threshold_capture(tmp_path)
    with pytest.raises(ValueError, match="outside its matrix cells"):
        script.build_track_m_threshold_capture(tmp_path, expected_sha256=None)
    (tmp_path / "thresh03.xlsx").unlink()
    with pytest.raises(FileNotFoundError, match="thresh03"):
        script.build_track_m_threshold_capture(tmp_path)
    (tmp_path / "thresh95.xlsx").unlink()
    with pytest.raises(FileNotFoundError, match="thresh95"):
        script.build_track_m_threshold_capture(tmp_path)


# -------------------------------------------------------------------------
# The loader and the coverage check (referee R8)
# -------------------------------------------------------------------------
def test_a_tampered_or_missing_capture_is_refused(tmp_path, capture):
    edited = copy.deepcopy(capture)
    edited["weighted_average"]["2010"]["one_65_plus"] += 1
    path = tmp_path / CAPTURE.name
    path.write_text(json.dumps(edited, indent=2) + "\n", encoding="utf-8")
    with pytest.raises(ValueError, match="sha256"):
        rules.load_aged_thresholds(path)
    with pytest.raises(rules.ThresholdsNotCapturedError, match="missing"):
        rules.load_aged_thresholds(tmp_path / "missing.json")


def test_a_missing_threshold_year_is_a_named_error():
    loaded = rules.load_aged_thresholds()
    # every year M4's structural count shows the in-window records need
    rules.check_threshold_years(
        {1982: 1, 1986: 1, 1988: 1, 1989: 1, 1991: 1, 1992: 1, 1998: 3},
        loaded,
    )
    rules.check_threshold_years({2003: 4, 2022: 1}, loaded)
    with pytest.raises(rules.ThresholdYearMissingError) as refused:
        rules.check_threshold_years(
            {1981: 1, 1987: 2, 1993: 1, 1996: 5, 2023: 1}, loaded
        )
    message = str(refused.value)
    assert (
        "threshold years 1981 (1 records), 1987 (2 records), 1993 (1 "
        "records), 2023 (1 records) are not in the capture" in message
    )
    assert "captured: 1982, 1986, 1988-1989, 1991-1992, 1994-2022" in message
    assert "1996" not in message
    assert "d279" in message
    for year in (1981, 1983, 1984, 1985, 1987, 1990, 1993, 2023):
        with pytest.raises(rules.ThresholdYearMissingError, match=str(year)):
            loaded.for_year(year)
    # a named error, not a KeyError
    assert not issubclass(rules.ThresholdYearMissingError, KeyError)
