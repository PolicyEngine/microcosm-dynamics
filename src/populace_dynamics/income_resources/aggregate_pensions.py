"""Pinned pension and household aggregates for D4's report-only WP8 inputs.

Invariants
----------
* Every exposed reader verifies its capture before inspecting table contents.
* Federal Reserve series retain positional dictionary line identities, so a
  repeated series code cannot overwrite another published table line.
* Workbook headers expand only published merged ranges; physical coordinates
  distinguish repeated row labels and worksheet continuations.
* Published missing/suppressed cells remain missing, never zero. Excel percent
  formats convert fractions to percent; monetary scales follow source headers.
* Readers return cells without printing/logging values; errors name structure.

The registry stores explicit workbook data regions and source column units.
These are aggregate publications, not Form 5500 or taxpayer record files.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from decimal import Decimal
from functools import lru_cache
from typing import TYPE_CHECKING

import openpyxl

from .aggregate_tables import (
    AggregateCell,
    ParsedNumber,
    TableSpec,
    TableStructureError,
    parse_number,
    verify_payload,
)

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet


def _text(value: object) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def parse_fed(spec: TableSpec, payload: bytes) -> tuple[AggregateCell, ...]:
    """Read one June-2026-code Z.1 table and its positional dictionary."""
    verify_payload(spec, payload)
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            csv_text = archive.read(spec.selector["csv_member"]).decode(
                "utf-8-sig"
            )
            dictionary = archive.read(
                spec.selector["dictionary_member"]
            ).decode("utf-8-sig")
    except (KeyError, ValueError, zipfile.BadZipFile, UnicodeDecodeError):
        raise TableStructureError("Z.1 archive structure mismatch") from None
    entries = list(csv.reader(io.StringIO(dictionary), delimiter="\t"))
    rows = csv.reader(io.StringIO(csv_text))
    header = next(rows, [])
    if not header or header[0] != "date" or len(header) != len(entries) + 1:
        raise TableStructureError("Z.1 dictionary/header width mismatch")
    for code, entry in zip(header[1:], entries, strict=True):
        if len(entry) != 5 or entry[0] != code:
            raise TableStructureError("Z.1 positional dictionary mismatch")
    result = []
    for row in rows:
        if len(row) != len(header):
            raise TableStructureError("Z.1 observation width mismatch")
        period = row[0]
        if not re.fullmatch(r"\d{4}(?:[: -]?Q[1-4])?", period):
            raise TableStructureError("Z.1 period label mismatch")
        year = int(period[:4])
        for entry, token in zip(entries, row[1:], strict=True):
            number = parse_number(token)
            result.append(
                AggregateCell(
                    spec.identifier,
                    " | ".join((entry[2], entry[1], entry[0])),
                    period,
                    year,
                    entry[4],
                    number.value,
                    number.footnote_flags,
                )
            )
    return _validate_cells(spec, result)


def _validate_cells(
    spec: TableSpec, cells: list[AggregateCell]
) -> tuple[AggregateCell, ...]:
    if not cells:
        raise TableStructureError("Registered table has no cells")
    if any(
        c.year not in spec.years or c.unit not in spec.units for c in cells
    ):
        raise TableStructureError("Cell coverage differs from registry")
    keys = [(c.row_key, c.column_key, c.year) for c in cells]
    if len(keys) != len(set(keys)):
        raise TableStructureError("Published cell keys are not distinct")
    return tuple(cells)


@lru_cache(maxsize=2)
def _workbook(payload: bytes):
    """Reuse a verified workbook across its many registered tables."""
    try:
        return openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
    except (ValueError, KeyError, zipfile.BadZipFile):
        raise TableStructureError("Workbook structure mismatch") from None


def _expanded_headers(sheet: Worksheet) -> dict[tuple[int, int], str]:
    headers = {
        (cell.row, cell.column): _text(cell.value)
        for row in sheet
        for cell in row
        if cell.value is not None
    }
    for merged in sheet.merged_cells.ranges:
        value = headers.get((merged.min_row, merged.min_col), "")
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                headers[row, column] = value
    return headers


def workbook_header(sheet: Worksheet, rows: list[int], column: int) -> str:
    """Join the distinct levels of one published merged column header."""
    expanded = _expanded_headers(sheet)
    labels = []
    for row in rows:
        label = expanded.get((row, column), "")
        if (
            label
            and label not in labels
            and not re.fullmatch(r"\(\d+\)", label)
        ):
            labels.append(label)
    return " / ".join(labels)


def _label_flags(label: str) -> tuple[str, ...]:
    flags = re.findall(r"\[([^\]]+)\]", label)
    # DOL puts ordinary trailing footnote digits in its Excel labels.
    flags.extend(re.findall(r"\s([1-9])(?=$|\s*/)", label))
    return tuple(flags)


def _parse_workbook(
    spec: TableSpec, payload: bytes
) -> tuple[AggregateCell, ...]:
    book = _workbook(payload)
    result = []
    for region in spec.selector["regions"]:
        name = region["sheet"]
        if name not in book.sheetnames:
            raise TableStructureError("Registered worksheet absent")
        sheet = book[name]
        expanded = _expanded_headers(sheet)
        columns = region["columns"]
        headers = {
            col: workbook_header(sheet, region["header_rows"], col)
            for col in columns
        }
        if any(not header for header in headers.values()):
            raise TableStructureError("Registered column header absent")
        for row in region["data_rows"]:
            labels = []
            for col in range(1, min(columns)):
                label = expanded.get((row, col), "")
                if label and label not in labels:
                    labels.append(label)
            if not labels:
                raise TableStructureError("Registered row label absent")
            row_key = f"{name}/r{row}: " + " / ".join(labels)
            for col in columns:
                cell = sheet.cell(row, col)
                token = _text(cell.value)
                # Bulletin notes define these as quantities below its
                # reporting threshold, not observed zeros.
                number = (
                    ParsedNumber(None, ("below_reporting_threshold", token))
                    if token in {"*/", "**/"}
                    else parse_number(token)
                )
                value = number.value
                unit = region["column_units"][str(col)]
                format_flags = tuple(
                    re.findall(r'"(\*+)\s*"', cell.number_format)
                )
                # IRS marks a hidden zero section as ** where the source
                # combines data to prevent disclosure. A numeric ** estimate
                # remains published and retains the combined-data flag.
                if value == 0 and re.search(r';;"\*\*"', cell.number_format):
                    value = None
                    format_flags += ("suppressed",)
                if value is not None and "%" in cell.number_format:
                    value *= Decimal(100)
                    if unit != "percent":
                        raise TableStructureError(
                            "Percent format/unit mismatch"
                        )
                year = region.get("column_years", {}).get(
                    str(col), region.get("year", 2023)
                )
                flags = tuple(
                    dict.fromkeys(
                        number.footnote_flags
                        + format_flags
                        + tuple(
                            f for label in labels for f in _label_flags(label)
                        )
                        + _label_flags(headers[col])
                    )
                )
                result.append(
                    AggregateCell(
                        spec.identifier,
                        row_key,
                        headers[col],
                        year,
                        unit,
                        value,
                        flags,
                    )
                )
    return _validate_cells(spec, result)


def parse_dol(spec: TableSpec, payload: bytes) -> tuple[AggregateCell, ...]:
    """Read one published bulletin table, including continuation sheets."""
    verify_payload(spec, payload)
    return _parse_workbook(spec, payload)


def parse_irs(spec: TableSpec, payload: bytes) -> tuple[AggregateCell, ...]:
    """Read one IRS IRA point-estimate table with count/amount units."""
    verify_payload(spec, payload)
    return _parse_workbook(spec, payload)
