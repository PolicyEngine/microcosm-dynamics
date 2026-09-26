"""Capture the parameters Tracks U and M read (plan items U3 and M2).

Two captures, each written to ``data/external`` and pinned by SHA-256 in
:mod:`populace_dynamics.estimates.adjusted_poverty`:

* ``--ssi`` writes ``track_u_ssi_parameters.json``: the federal SSI
  benefit rate (individual and couple, monthly) in force on January 1 of
  each income year 2004-2012, the $20 general and $65 earned income
  exclusions, the share of remaining earned income excluded, and the
  individual and couple resource limits.  The values come from the
  policyengine-us parameter files (``gov/ssa/ssi``) of the checkout given
  by ``--pe-us-dir`` (default ``POPULACE_DYNAMICS_PE_US_DIR`` or
  ``~/PolicyEngine/policyengine-us``); the capture records each file's
  SHA-256, the checkout revision and each file's cited reference.  These
  are assumptions for the SSI response rule, not comparator values.
* ``--census-dir DIR`` writes ``census_poverty_thresholds_2004_2012.json``
  from the Census Bureau's threshold workbooks for income years 2004-2012
  staged in ``DIR`` as ``thresh04.xlsx`` ... ``thresh12.xlsx`` (the file
  names the Census historical thresholds page lists under
  ``https://www2.census.gov/programs-surveys/cps/tables/time-series/
  historical-poverty-thresholds/``).  The nine workbooks were staged on
  2026-09-24 under cos decision d194 and are committed in
  ``data/external/census_poverty_thresholds/``; each one's SHA-256 is
  pinned in :data:`CENSUS_WORKBOOK_SHA256` and checked before it is
  parsed.  The parser was written against the real layout (inspected cell
  by cell in every workbook, ``thresh03.xlsx`` included), locates rows
  and columns by their printed labels and refuses any other layout, and
  checks every year's values within the year and against the previous
  year (:func:`parse_threshold_rows`, :func:`check_matrix_moves_together`).
  Re-running it on the committed workbooks reproduces the committed
  capture byte for byte (``tests/track_u/test_census_threshold_capture.py``).
* ``--track-m-census-dir DIR`` writes ``census_poverty_thresholds_1982_2022.json``
  for Track M (DynaSim exercise 4, the minimum benefit) from the
  thirty-five workbooks of 1982, 1986, 1988, 1989, 1991, 1992 and
  1994-2022: the nine above, ``thresh03.xlsx`` (staged under d194 with
  them), ``thresh13.xlsx`` ... ``thresh22.xlsx`` (fetched on 2026-09-25 by
  the orchestrating Claude Code session after Max approved the download in
  cos decision d279) and the fifteen years before 2003 that M4's
  structural count shows the in-window records need (staged on 2026-09-25
  under d279's "any earlier year the build proves it needs";
  ``thresh95.xlsx`` from the Internet Archive's copy of its Census URL,
  :data:`ARCHIVE_RETRIEVALS`).  All thirty-five are committed in
  ``data/external/census_poverty_thresholds/`` and pinned in
  :data:`CENSUS_WORKBOOK_SHA256`.  The same parser reads them.  Inspected
  cell by cell (2026-09-25), every workbook prints the 2003-2012 table in
  the same cells with the same labels; the parser accepts exactly these
  departures, each pinned to its years: ``thresh19.xlsx`` carries two
  further worksheets, ``Sheet2`` and ``Sheet3``, that must be empty
  (:data:`EMPTY_EXTRA_SHEETS`); ``thresh22.xlsx`` prints every weighted
  average rounded to $10 while its matrix cells stay whole dollars
  (:data:`WEIGHTED_AVERAGE_UNIT`); the notes of 1982-2000 name the March
  CPS and 2001's the CPS ADS (:data:`NOTE_SURVEY_BEFORE_2002`); 2001's
  size rows say "persons" (:data:`PERSONS_ROW_LABEL_YEARS`); and 1982 and
  2000 print one revision line below the note (:data:`REVISION_LINES`).
  Track M reads the weighted average for one person aged 65 and over
  (``one_65_plus``) as printed; the pin is
  ``min_benefit_track_m.thresholds.TRACK_M_THRESHOLDS_SHA256``, and
  ``tests/min_benefit_track_m/test_threshold_capture.py`` and
  ``test_threshold_capture_before_2003.py`` re-run the capture and compare
  it with the workbook cells and with Census's HTML Table 1.

Usage::

    python scripts/capture_track_u_parameters.py --ssi [--pe-us-dir DIR]
    python scripts/capture_track_u_parameters.py --census-dir DIR
    python scripts/capture_track_u_parameters.py --track-m-census-dir DIR
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

import yaml

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

from populace_dynamics.estimates import adjusted_poverty  # noqa: E402
from populace_dynamics.min_benefit_track_m import (  # noqa: E402
    thresholds as track_m_thresholds,
)

#: Income years the exact-age primary and the pooled row observe.
YEARS: tuple[int, ...] = tuple(range(2004, 2013))
_PE_US_ENV = "POPULACE_DYNAMICS_PE_US_DIR"
_PE_US_DEFAULT = Path("~/PolicyEngine/policyengine-us").expanduser()
_SSI_ROOT = Path("policyengine_us/parameters/gov/ssa/ssi")
SSI_FILES: dict[str, Path] = {
    "fbr_individual": _SSI_ROOT / "amount" / "individual.yaml",
    "fbr_couple": _SSI_ROOT / "amount" / "couple.yaml",
    "general_income_exclusion": _SSI_ROOT
    / "income"
    / "exclusions"
    / "general.yaml",
    "earned_income_exclusion": _SSI_ROOT
    / "income"
    / "exclusions"
    / "earned.yaml",
    "earned_income_share_excluded": _SSI_ROOT
    / "income"
    / "exclusions"
    / "earned_share.yaml",
    "resource_limit_individual": _SSI_ROOT
    / "eligibility"
    / "resources"
    / "limit"
    / "individual.yaml",
    "resource_limit_couple": _SSI_ROOT
    / "eligibility"
    / "resources"
    / "limit"
    / "couple.yaml",
}
CENSUS_URL_BASE = (
    "https://www2.census.gov/programs-surveys/cps/tables/time-series/"
    "historical-poverty-thresholds/"
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def resolve_pe_us(pe_us_dir: Path | None) -> Path:
    if pe_us_dir is not None:
        return Path(pe_us_dir).expanduser()
    env = os.environ.get(_PE_US_ENV)
    return Path(env).expanduser() if env else _PE_US_DEFAULT


def _in_force(values: dict[Any, Any], year: int) -> float:
    """The value of a dated parameter in force on January 1 of ``year``."""

    dated = sorted(
        (dt.date.fromisoformat(str(key)), float(value))
        for key, value in values.items()
    )
    start = dt.date(year, 1, 1)
    eligible = [value for date, value in dated if date <= start]
    if not eligible:
        raise ValueError(f"no value in force on {start}")
    return eligible[-1]


def build_ssi_capture(pe_us_dir: Path | None = None) -> dict[str, Any]:
    root = resolve_pe_us(pe_us_dir).resolve()
    documents: dict[str, dict[str, Any]] = {}
    hashes: dict[str, str] = {}
    references: dict[str, list[dict[str, str]]] = {}
    for name, relative in SSI_FILES.items():
        path = root / relative
        if not path.is_file():
            raise FileNotFoundError(f"policyengine-us file missing: {path}")
        hashes[str(relative)] = _sha256(path)
        documents[name] = yaml.safe_load(path.read_text())
        references[name] = [
            {"title": ref.get("title", ""), "href": ref.get("href", "")}
            for ref in documents[name].get("metadata", {}).get("reference", [])
        ]
    try:
        revision = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=root,
            capture_output=True,
            text=True,
            check=True,
        ).stdout.strip()
    except (subprocess.CalledProcessError, FileNotFoundError):
        revision = "unknown"

    def year_map(name: str) -> dict[str, float]:
        return {
            str(year): _in_force(documents[name]["values"], year)
            for year in YEARS
        }

    def constant(name: str) -> float:
        values = {_in_force(documents[name]["values"], year) for year in YEARS}
        if len(values) != 1:
            raise ValueError(f"{name} changes within {YEARS[0]}-{YEARS[-1]}")
        return values.pop()

    return {
        "schema_version": adjusted_poverty.SSI_SCHEMA_VERSION,
        "description": (
            "Federal SSI parameters for the Track U SSI response rules "
            "(exercise 2; Python income concept, not Axiom): the federal "
            "benefit rate in force on January 1 of each income year, the "
            "general and earned income exclusions and the resource limits"
        ),
        "years": [YEARS[0], YEARS[-1]],
        "source": {
            "policyengine_us_revision": revision,
            "policyengine_us_files_sha256": hashes,
            "references": references,
            "rule": "value in force on January 1 of the income year",
            "generated_by": "scripts/capture_track_u_parameters.py --ssi",
        },
        "federal_benefit_rate_monthly": {
            "individual": year_map("fbr_individual"),
            "couple": year_map("fbr_couple"),
        },
        "general_income_exclusion_monthly": constant(
            "general_income_exclusion"
        ),
        "earned_income_exclusion_monthly": constant("earned_income_exclusion"),
        "earned_income_share_excluded": constant(
            "earned_income_share_excluded"
        ),
        "resource_limit": {
            "individual": constant("resource_limit_individual"),
            "couple": constant("resource_limit_couple"),
        },
    }


# ---------------------------------------------------------------------------
# Census thresholds
# ---------------------------------------------------------------------------
#: The committed copies of the Census workbooks the capture reads.
CENSUS_WORKBOOK_DIR = ROOT / "data" / "external" / "census_poverty_thresholds"
#: SHA-256 of each Census workbook the captures read, from
#: :data:`CENSUS_URL_BASE` and committed in :data:`CENSUS_WORKBOOK_DIR`:
#: ``thresh03.xlsx`` ... ``thresh12.xlsx`` fetched on 2026-09-24 under cos
#: decision d194 and ``thresh13.xlsx`` ... ``thresh22.xlsx`` on 2026-09-25
#: under cos decision d279, each by the orchestrating Claude Code session
#: after Max approved the download, and the fifteen years before 2003
#: below (the staging directory's ``SHA256SUMS`` lists all thirty-five
#: digests as of 2026-09-25).  A workbook with other bytes is refused
#: before it is parsed.
CENSUS_WORKBOOK_SHA256: dict[str, str] = {
    "thresh03.xlsx": (
        "f91f2a70062c52b21391e74dc0486bf9adc898c8c923212ecaa7db7a0db24895"
    ),
    "thresh04.xlsx": (
        "9cae1bcff6c3ff80faedaa11c28068d9a640f10f5a7679e3cfb0430779bbb5b8"
    ),
    "thresh05.xlsx": (
        "9c626a9757232d500167349c9c47dc254a5f1ac34c6ba3468acd1a624df52888"
    ),
    "thresh06.xlsx": (
        "7ace9e1c3990348b252698da9026aea082c7a52bcc073c0315721b01b9c09a45"
    ),
    "thresh07.xlsx": (
        "1209903a9c1071f9a76ee5a50dc9eb83faa93024134a72e73a3cdc6b558a4327"
    ),
    "thresh08.xlsx": (
        "7e7222d431411aa82734d9580cc8f6bce246d1cd330e688a8306748bc46f09ae"
    ),
    "thresh09.xlsx": (
        "7f997506d189bb1899bab1255ae5e991a384ffaf8d72d63ca4497d1912e1175e"
    ),
    "thresh10.xlsx": (
        "75360b3d852669c76f77df20d2292d1abe1d90e1b58a9d3e0611bb37570a45b9"
    ),
    "thresh11.xlsx": (
        "ec2758efc4b6797b171fbe4cda135e5f08396734bdc2d606f0b990219f42bab6"
    ),
    "thresh12.xlsx": (
        "26da2dc48de0a0799905c15aa6634d5b942dd57b801bdaf9564eacf4cc365f36"
    ),
    "thresh13.xlsx": (
        "4114b7a98043842fc6ec9cc688d6918ce1461f84ed5fc96ffe0c8a3076083f2b"
    ),
    "thresh14.xlsx": (
        "c9fad68272d238036bf01e24ea33c1de97f71c6d56b310276f173b5b21153ce8"
    ),
    "thresh15.xlsx": (
        "9658e7c7fa63fe25a396f1abfcc9c90db385edbeb19ffdabf5e6d14eaf650e41"
    ),
    "thresh16.xlsx": (
        "5d16803e3904430564b7df980b7968563634a9e0619b96bcbc079b1dcb854e45"
    ),
    "thresh17.xlsx": (
        "edf7af0544b48cd8de86cfd12d018c5dcc596677cc9f532c070eb18262db3e65"
    ),
    "thresh18.xlsx": (
        "f1abec2ee137a39e04466ec5cf189412228d2c1e5a2d56fd5195bb90ca6ffa55"
    ),
    "thresh19.xlsx": (
        "e9252e05ef17d0787243eeadec1228524170211f57ac3b58efe32ee909e8ff57"
    ),
    "thresh20.xlsx": (
        "5739e473550312b7663479711f41d254b477a7d020d92847cb6164dc506767f8"
    ),
    "thresh21.xlsx": (
        "9399f4564ed22776f286fbceb72bab288f0cf0c1a70181f3805b3321332734e5"
    ),
    "thresh22.xlsx": (
        "5874eb8ecc525f5d26daab34b81165416c669daf62771e65ef52847bd3cbf89f"
    ),
    # The fifteen years before 2003 that M4's structural count shows the
    # in-window records need (1982, 1986, 1988, 1989, 1991, 1992 and
    # 1994-2002), staged on 2026-09-25 under cos decision d279 ("any
    # earlier year the build proves it needs") in the directory whose
    # SHA256SUMS lists all thirty-five digests: fourteen from
    # CENSUS_URL_BASE and thresh95.xlsx from the Internet Archive's copy
    # of its Census URL (ARCHIVE_RETRIEVALS).
    "thresh82.xlsx": (
        "88a5fb3aaba6008e430bce327b862fca03ba85340c32c1375b902dfb60093b78"
    ),
    "thresh86.xlsx": (
        "49c00638209721b4cf25e4039728255b97874efd54f36ff085860113b2d077b2"
    ),
    "thresh88.xlsx": (
        "f22fc66ea80677c34bf36d00f9c26aff3038a0f115f90673a1ebd42c7f19fd11"
    ),
    "thresh89.xlsx": (
        "bed78af3dd2425349f33391163cc76a1218080403eeaa6e7c0de15d359847d73"
    ),
    "thresh91.xlsx": (
        "a06989525bc712f609fa40cde808ec1cfb5bf8ce1d8194eeb362f2ec7847249c"
    ),
    "thresh92.xlsx": (
        "ac5262b83d0f63cff97caa5ad8fb80096f8b16db501f79dd88cca19e9b4fd3d5"
    ),
    "thresh94.xlsx": (
        "593b5a45109ed58ea35648e0f9c6ba749050741e4e935ab69cce9bbc15d38006"
    ),
    "thresh95.xlsx": (
        "32bb678f3e0847b71c96c7c3c2b2ca8151a5c82307466df2807e9438ab2138f8"
    ),
    "thresh96.xlsx": (
        "28a0cb77b312f861dd55b51d68cf83452d8a17c93f7b111e1b6da9930e27db78"
    ),
    "thresh97.xlsx": (
        "aa463ec2d30e7d89a8626cd0ea73caeb15e2351b566e20163625dff0d8389754"
    ),
    "thresh98.xlsx": (
        "c8ab5b79a49a756952a31c684a961f5da7de8dce46ce60c50128c618618964bb"
    ),
    "thresh99.xlsx": (
        "a513a440417c603a7db00dbec785cb5c23851bb156bc0bfc4fc28e460a2f3a63"
    ),
    "thresh00.xlsx": (
        "46a08f5146c5cfdb6cc6c6fb9e2a50deb78e68acbfa5917fbd5bfaa2d211f3bb"
    ),
    "thresh01.xlsx": (
        "ac7deb00d0feee4fe0ca8124ecb19ef85d87bf11fdc39ac608d078b5bfe468d2"
    ),
    "thresh02.xlsx": (
        "40609e1d94d2e1972cfbab6904ebee9e03d94686eaf9af54dafe782d4d7f648d"
    ),
}
#: Workbooks not fetched from their Census URL, and how they were, as the
#: staging directory's ``PROVENANCE-thresh95.md`` records it (copied into
#: ``census_poverty_thresholds/provenance.md``).  On 2026-09-25 the Census
#: server answered the thresh95.xlsx URL, and only that one, with its WAF
#: page "Request Rejected" (HTTP 200, 247 bytes of HTML).  The file is the
#: Internet Archive's raw (``id_``) copy of that URL, fetched from two
#: captures 3.5 years apart that are byte-identical once the 2023 one is
#: decompressed, each matching the SHA-1 digest the Archive's CDX index
#: records for it.
ARCHIVE_RETRIEVALS: dict[str, dict[str, Any]] = {
    "thresh95.xlsx": {
        "retrieved_from": "internet_archive",
        "reason": (
            "the Census server answered this URL, and only this one, with "
            "its WAF page 'Request Rejected' (HTTP 200, 247 bytes of HTML) "
            "on 2026-09-25"
        ),
        "archive_captures": [
            {
                "url": (
                    "http://web.archive.org/web/20230226205734id_/"
                    + CENSUS_URL_BASE
                    + "thresh95.xlsx"
                ),
                "served": "gzip-encoded; decompressed",
                "cdx_sha1_base32": "ARBZSSCAOHTKTGHH6INRSDDWBDJVFVAB",
                "cdx_sha1_of": "the gzip body as served",
            },
            {
                "url": (
                    "http://web.archive.org/web/20260820215832id_/"
                    + CENSUS_URL_BASE
                    + "thresh95.xlsx"
                ),
                "served": "the file itself",
                "cdx_sha1_base32": "ZCDVY76QW2OH6NA6RLKWJPMJCNTMZIAQ",
                "cdx_sha1_of": "the file",
            },
        ],
        "captures_identical": True,
        "to_do": (
            "retry the Census URL when its WAF allows it and confirm the "
            "SHA-256"
        ),
    },
}
#: Worksheets a workbook may carry beyond the table's, by income year; each
#: must be empty.  Only ``thresh19.xlsx`` has any (inspected 2026-09-25):
#: ``Sheet2`` and ``Sheet3``, with no cell value.
EMPTY_EXTRA_SHEETS: dict[int, tuple[str, ...]] = {2019: ("Sheet2", "Sheet3")}
#: The survey each year's note names as the source of its weighted
#: averages, for the years before 2002 (inspected cell by cell,
#: 2026-09-25): 1982-2000 print "the March <year + 1> Current Population
#: Survey (CPS)", 2001 "the 2002 Current Population Survey Annual
#: Demographic Supplement (CPS ADS)".  From 2002 every note names the
#: <year + 1> CPS ASEC.  A year before 2002 that is not listed has no
#: inspected layout and is refused.
NOTE_SURVEY_BEFORE_2002: dict[int, str] = {
    **{
        year: "march_cps"
        for year in (1982, 1986, 1988, 1989, 1991, 1992, *range(1994, 2001))
    },
    2001: "cps_ads",
}
#: Income years whose size rows read "Two persons" ... "Nine persons or
#: more" where every other year prints "people" (thresh01.xlsx only,
#: inspected 2026-09-25).
PERSONS_ROW_LABEL_YEARS: frozenset[int] = frozenset({2001})
#: The one line a workbook may print below its note, by income year, as
#: printed (whitespace collapsed; inspected 2026-09-25).  Census revised
#: thresh82.xlsx and thresh00.xlsx after first publishing them; no other
#: year carries such a line.
REVISION_LINES: dict[int, str] = {
    1982: "Revised on 4/19/2022 due to rounding issues.",
    2000: "Revised on 2/1/2023 due to a formatting error.",
}
#: The dollar unit the weighted averages are printed to, by income year
#: (whole dollars unless listed).  Every weighted average in thresh22.xlsx
#: is a multiple of $10 and its matrix cells are whole dollars, so the
#: size-1 age rows print 15,230 and 14,040 beside single matrix cells of
#: 15,225 and 14,036 (inspected 2026-09-25).  For such a year every
#: weighted average must be a multiple of the unit, and it may lie up to
#: half the unit outside its row's matrix cells; in every other year the
#: weighted average must lie within them exactly.
WEIGHTED_AVERAGE_UNIT: dict[int, int] = {2022: 10}

#: The caption every real workbook prints in A1, above the title (lower
#: case, whitespace collapsed).
_CAPTION = (
    "table with row headings in column a and column headings in rows 5 to 6."
)
#: The table title, matched in full (lower case, whitespace collapsed).
_TITLE = re.compile(
    r"poverty thresholds for (\d{4}) by size of family and number of "
    r"related children under 18 years"
)
_UNITS = "(in dollars)"
_SOURCE = re.compile(r"source: u\.s\. census bureau, (\d{4})\.?")
#: The note's first sentence, by the survey it names (the 2009 note adds
#: two more).  Which survey each year's note must name is
#: :func:`note_survey`'s.
_NOTE_PATTERNS: dict[str, re.Pattern[str]] = {
    "cps_asec": re.compile(
        r"note: the source of the weighted average thresholds is the "
        r"(\d{4}) current population survey annual social and economic "
        r"supplement \(cps asec\)\."
    ),
    "cps_ads": re.compile(
        r"note: the source of the weighted average thresholds is the "
        r"(\d{4}) current population survey annual demographic supplement "
        r"\(cps ads\)\."
    ),
    "march_cps": re.compile(
        r"note: the source of the weighted average thresholds is the march "
        r"(\d{4}) current population survey \(cps\)\."
    ),
}
_NOTE = _NOTE_PATTERNS["cps_asec"]
#: How the refusals name each survey.
_SURVEY_NAMES = {
    "cps_asec": "CPS ASEC",
    "cps_ads": "CPS ADS",
    "march_cps": "March CPS",
}
#: The only text the note may carry after its first sentence: the two
#: sentences of the 2009 note (a year whose average annual CPI-U fell),
#: with ``{year}`` the table's year and ``{previous}`` the year before.
_NOTE_CPI_FALL = (
    "the poverty thresholds are updated each year using the change in the "
    "average annual consumer price index for all urban consumers (cpi-u). "
    "since the average annual cpi-u for {year} was lower than the average "
    "annual cpi-u for {previous}, poverty thresholds for {year} are "
    "slightly lower than the corresponding thresholds for {previous}."
)
#: The table's labelled rows in printed order: (label, key, number of
#: related-children columns the row fills).  ``one_all`` and ``two_all``
#: are the size-1 and size-2 rows over both age groups, which print a
#: weighted average only.
_TABLE_ROWS: tuple[tuple[str, str, int], ...] = (
    ("one person (unrelated individual)", "one_all", 0),
    ("under 65 years", "one_under_65", 1),
    ("65 years and over", "one_65_plus", 1),
    ("two people", "two_all", 0),
    ("householder under 65 years", "two_under_65", 2),
    ("householder 65 years and over", "two_65_plus", 2),
    ("three people", "three", 3),
    ("four people", "four", 4),
    ("five people", "five", 5),
    ("six people", "six", 6),
    ("seven people", "seven", 7),
    ("eight people", "eight", 8),
    ("nine people or more", "nine_plus", 9),
)
_ALL_AGES = {"one_all": "one", "two_all": "two"}


def note_survey(year: int) -> str:
    """The survey ``year``'s note must name (a key of the note patterns).

    The CPS ASEC from 2002; before that, the year's inspected layout
    (:data:`NOTE_SURVEY_BEFORE_2002`).  A year before 2002 with no inspected
    layout is refused.
    """

    if year >= 2002:
        return "cps_asec"
    if year not in NOTE_SURVEY_BEFORE_2002:
        raise ValueError(
            f"{year}: no inspected layout for a Census threshold workbook "
            "before 2002 (NOTE_SURVEY_BEFORE_2002 lists "
            f"{sorted(NOTE_SURVEY_BEFORE_2002)})"
        )
    return NOTE_SURVEY_BEFORE_2002[year]


def row_labels(year: int) -> list[str]:
    """The thirteen row labels ``year``'s table prints, in order.

    Every year prints :data:`_TABLE_ROWS`' labels, except that a year in
    :data:`PERSONS_ROW_LABEL_YEARS` prints "persons" for "people" in the
    size rows ("Two persons" ... "Nine persons or more").
    """

    labels = [label for label, _, _ in _TABLE_ROWS]
    if year in PERSONS_ROW_LABEL_YEARS:
        labels = [
            re.sub(r"(?<=\w) people\b", " persons", label) for label in labels
        ]
    return labels


#: The related-children column headers, in printed order (0 ... 8).
_CHILD_HEADERS: tuple[str, ...] = (
    "none",
    "one",
    "two",
    "three",
    "four",
    "five",
    "six",
    "seven",
    "eight or more",
)
#: Columns: row label, weighted average, then the nine children columns.
_LABEL_COL, _WEIGHTED_COL, _FIRST_CHILD_COL = 0, 1, 2
_TABLE_WIDTH = _FIRST_CHILD_COL + len(_CHILD_HEADERS)
#: Largest distance, in dollars, of a matrix cell from the previous
#: year's cell times the year's common ratio (whole-dollar rounding of
#: both years moves a cell by at most about $1).
_CROSS_YEAR_TOLERANCE = 2.0


def _text(value: Any) -> str:
    return " ".join(str(value).split()).lower()


def _label(value: Any) -> str:
    """A printed row or column label without its dot leaders or colon."""

    return re.sub(r"[\s.:…]+$", "", _text("" if value is None else value))


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _dollars(value: Any, where: str) -> int:
    """A threshold cell: a positive whole number of dollars, or refused."""

    if isinstance(value, bool) or not isinstance(value, int | float):
        raise ValueError(f"{where}: {value!r} is not a number")
    if not float(value).is_integer() or value <= 0:
        raise ValueError(f"{where}: {value!r} is not a positive whole dollar")
    return int(value)


def parse_threshold_rows(rows: list[list[Any]], year: int) -> dict[str, Any]:
    """Parse one year's Census threshold table from its cell rows.

    The layout is the one every ``thresh03.xlsx`` ... ``thresh12.xlsx``
    shares (checked on the real workbooks, 2026-09-24): above the header
    exactly three lines in column A, the caption :data:`_CAPTION`, the
    title and "(In dollars)"; a header row "Size of family unit" |
    "Weighted average thresholds" | "Related children under 18 years"
    over a row of the nine children headers "None" ... "Eight or more";
    the thirteen labelled rows of :data:`_TABLE_ROWS` in order; then
    "Source: U.S. Census Bureau, <year + 1>." and a note naming the
    <year + 1> CPS ASEC, which may add only the two CPI-U sentences of
    :data:`_NOTE_CPI_FALL` (the 2009 note).  Everything else must be
    blank (whitespace-only cells count as blank).

    The workbooks before 2003 (inspected cell by cell, 2026-09-25) print
    the same table in the same cells, with three departures, each pinned
    to its years: the note names the March <year + 1> CPS (1982-2000) or
    the <year + 1> CPS ADS (2001) instead of the CPS ASEC
    (:func:`note_survey`); 2001's size rows say "persons" for "people"
    (:func:`row_labels`); and 1982 and 2000 print one revision line below
    the note (:data:`REVISION_LINES`), which must be exactly that line.
    Header capitalization and line breaks ("Weighted Average Thresholds",
    "Eight or more" over three lines) do not matter: labels are compared
    in lower case with whitespace collapsed.

    Refuses a table whose title does not name ``year``, whose source or
    note does not name ``year + 1`` (and, in the note, the year's survey),
    with other text above the header or in the note, with a line below
    the note other than its year's revision line, whose header or row
    labels differ, whose threshold cells are not positive whole numbers
    or fill other children columns than their row allows, with a value
    anywhere else, or whose values fail :func:`_validate_year`.

    Returns ``{"weighted_average": {row_key: dollars}, "all_ages":
    {"one": dollars, "two": dollars}, "matrix": {row_key: {children:
    dollars}}, "title": str, "source_line": str, "survey": str,
    "survey_year": int, "revision_line": str | None}``, plus
    ``"cps_asec_year"`` (the survey year) when the survey is the CPS
    ASEC.
    """

    width = max([_TABLE_WIDTH, *(len(row) for row in rows)])
    grid = [list(row) + [None] * (width - len(row)) for row in rows]

    def refuse(message: str) -> ValueError:
        return ValueError(f"{year}: {message}")

    titles = [
        (i, cell, match)
        for i, row in enumerate(grid)
        for cell in row
        if isinstance(cell, str) and (match := _TITLE.fullmatch(_text(cell)))
    ]
    if len(titles) != 1:
        raise refuse(
            "expected one 'Poverty Thresholds for <year> by Size of Family "
            "and Number of Related Children Under 18 Years' title, found "
            f"{len(titles)}"
        )
    title_row, title_cell, title_match = titles[0]
    if int(title_match.group(1)) != year:
        raise refuse(f"the title names {title_match.group(1)}")
    header = [
        i
        for i, row in enumerate(grid)
        if _label(row[_LABEL_COL]) == "size of family unit"
    ]
    if len(header) != 1:
        raise refuse("expected one 'Size of family unit' header row")
    h = header[0]
    if title_row >= h:
        raise refuse("the title is not above the header")
    above = [
        (column, cell)
        for row in grid[:h]
        for column, cell in enumerate(row)
        if not _blank(cell)
    ]
    if any(not isinstance(cell, str) for _, cell in above):
        raise refuse("a value above the header row")
    if any(column != _LABEL_COL for column, _ in above):
        raise refuse("text above the header outside column A")
    # Exactly the caption, the title and the units line, in that order
    # (each is a cell of its own in column A): no other text, formula or
    # line above the header.
    lines = [_text(cell) for _, cell in above]
    if (
        len(lines) != 3
        or lines[0] != _CAPTION
        or _TITLE.fullmatch(lines[1]) is None
        or lines[2] != _UNITS
    ):
        raise refuse(
            "the text above the header is not exactly the caption, the "
            "title and '(In dollars)', in that order"
        )
    top = grid[h]
    if _label(top[_WEIGHTED_COL]) != "weighted average thresholds":
        raise refuse("column B of the header is not 'Weighted average'")
    if _label(top[_FIRST_CHILD_COL]) != "related children under 18 years":
        raise refuse("column C of the header is not 'Related children'")
    if any(not _blank(cell) for cell in top[_FIRST_CHILD_COL + 1 :]):
        raise refuse("unexpected cells in the header row")
    if h + 1 >= len(grid):
        raise refuse("no children header row")
    children = grid[h + 1]
    printed = [
        _label(cell) for cell in children[_FIRST_CHILD_COL:_TABLE_WIDTH]
    ]
    if (
        printed != list(_CHILD_HEADERS)
        or not all(_blank(cell) for cell in children[:_FIRST_CHILD_COL])
        or not all(_blank(cell) for cell in children[_TABLE_WIDTH:])
    ):
        raise refuse(f"children headers {printed} != {list(_CHILD_HEADERS)}")
    sources = [
        i
        for i, row in enumerate(grid)
        if isinstance(row[_LABEL_COL], str)
        and _text(row[_LABEL_COL]).startswith("source:")
    ]
    if len(sources) != 1 or sources[0] <= h + 1:
        raise refuse("expected one 'Source:' line below the table")
    s = sources[0]
    source_line = " ".join(str(grid[s][_LABEL_COL]).split())
    source_match = _SOURCE.fullmatch(_text(source_line))
    if source_match is None or int(source_match.group(1)) != year + 1:
        raise refuse(
            f"source line {source_line!r} is not 'Source: U.S. Census "
            f"Bureau, {year + 1}.'"
        )
    survey = note_survey(year)
    note = grid[s + 1][_LABEL_COL] if s + 1 < len(grid) else None
    note_match = (
        None if _blank(note) else _NOTE_PATTERNS[survey].match(_text(note))
    )
    if note_match is None or int(note_match.group(1)) != year + 1:
        raise refuse(
            f"the note does not name the {year + 1} {_SURVEY_NAMES[survey]}"
        )
    tail = _text(note)[note_match.end() :].strip()
    if tail not in ("", _NOTE_CPI_FALL.format(year=year, previous=year - 1)):
        raise refuse(
            f"the note says more than the {_SURVEY_NAMES[survey]} sentence "
            "(and, for a year whose CPI-U fell, the two CPI-U sentences): "
            f"{tail!r}"
        )
    trailer = [grid[s][_LABEL_COL + 1 :], grid[s + 1][_LABEL_COL + 1 :]]
    below = grid[s + 2 :]
    revision_line = None
    if year in REVISION_LINES:
        printed = below[0][_LABEL_COL] if below else None
        if _blank(printed) or _text(printed) != _text(REVISION_LINES[year]):
            raise refuse(
                f"expected the revision line {REVISION_LINES[year]!r} "
                f"below the note, found {printed!r}"
            )
        revision_line = " ".join(str(printed).split())
        trailer.append(below[0][_LABEL_COL + 1 :])
        below = below[1:]
    trailer.extend(below)
    if any(not _blank(cell) for row in trailer for cell in row):
        raise refuse("a value beside or below the source and note")
    labelled = []
    for row in grid[h + 2 : s]:
        if _blank(row[_LABEL_COL]):
            if any(not _blank(cell) for cell in row):
                raise refuse("a value in an unlabelled row")
            continue
        labelled.append(row)
    labels = [_label(row[_LABEL_COL]) for row in labelled]
    expected = row_labels(year)
    if labels != expected:
        raise refuse(f"row labels {labels} != {expected}")
    weighted: dict[str, int] = {}
    all_ages: dict[str, int] = {}
    matrix: dict[str, dict[int, int]] = {}
    for (label, key, n_columns), row in zip(
        _TABLE_ROWS, labelled, strict=True
    ):
        where = f"{year} row {label!r}"
        value = _dollars(row[_WEIGHTED_COL], f"{where} weighted average")
        cells = {
            k: _dollars(cell, f"{where} column {_CHILD_HEADERS[k]!r}")
            for k, cell in enumerate(row[_FIRST_CHILD_COL:_TABLE_WIDTH])
            if not _blank(cell)
        }
        if sorted(cells) != list(range(n_columns)):
            raise refuse(
                f"row {label!r} fills children columns {sorted(cells)}, "
                f"expected {list(range(n_columns))}"
            )
        if any(not _blank(cell) for cell in row[_TABLE_WIDTH:]):
            raise refuse(f"row {label!r} has a value beyond the table")
        if key in _ALL_AGES:
            all_ages[_ALL_AGES[key]] = value
        else:
            weighted[key] = value
            matrix[key] = cells
    _validate_year(year, weighted, all_ages, matrix)
    parsed = {
        "weighted_average": weighted,
        "all_ages": all_ages,
        "matrix": matrix,
        "title": " ".join(str(title_cell).split()),
        "source_line": source_line,
        "survey": survey,
        "survey_year": int(note_match.group(1)),
        "revision_line": revision_line,
    }
    if survey == "cps_asec":
        parsed["cps_asec_year"] = parsed["survey_year"]
    return parsed


def _validate_year(
    year: int,
    weighted: dict[str, int],
    all_ages: dict[str, int],
    matrix: dict[str, dict[int, int]],
) -> None:
    """Relations every Census threshold table satisfies, or refused.

    The 65-and-over rows lie below the under-65 rows (sizes 1 and 2); the
    size-1 and size-2 weighted averages over both ages lie between their
    two age rows; the weighted averages rise with size from two persons;
    and each row's weighted average lies within the range of that row's
    matrix cells (a weighted average of them), which for sizes 1 ties the
    two columns together.  A year listed in :data:`WEIGHTED_AVERAGE_UNIT`
    prints its weighted averages rounded to that unit: each must be a
    multiple of it and may lie up to half of it outside the range.
    """

    if set(weighted) != set(adjusted_poverty.THRESHOLD_ROW_KEYS):
        raise ValueError(f"{year}: rows {sorted(weighted)}")
    unit = WEIGHTED_AVERAGE_UNIT.get(year, 1)
    if unit > 1:
        off_unit = sorted(
            key
            for key, value in (*weighted.items(), *all_ages.items())
            if value % unit
        )
        if off_unit:
            raise ValueError(
                f"{year}: weighted averages {off_unit} are not multiples of "
                f"${unit}, the unit this year's workbook prints them to"
            )
    slack = unit / 2 if unit > 1 else 0
    for size in ("one", "two"):
        young, old = weighted[f"{size}_under_65"], weighted[f"{size}_65_plus"]
        if not old < young:
            raise ValueError(
                f"{year}: 65-and-over threshold not below under-65 ({size})"
            )
        if not old <= all_ages[size] <= young:
            raise ValueError(
                f"{year}: the {size}-person weighted average over both ages "
                "is not between its two age rows"
            )
    larger = ["three", "four", "five", "six", "seven", "eight", "nine_plus"]
    chain = [weighted["two_under_65"], *(weighted[key] for key in larger)]
    if any(b <= a for a, b in zip(chain, chain[1:], strict=False)):
        raise ValueError(f"{year}: weighted averages not increasing in size")
    for key, value in weighted.items():
        cells = matrix[key].values()
        if not min(cells) - slack <= value <= max(cells) + slack:
            raise ValueError(
                f"{year}: row {key} weighted average {value} outside its "
                f"matrix cells {min(cells)}-{max(cells)}"
                + (
                    f" (+/- ${slack:g} for its ${unit} rounding)"
                    if slack
                    else ""
                )
            )


def check_matrix_moves_together(
    earlier: dict[str, dict[int, int]],
    later: dict[str, dict[int, int]],
    year: int,
    earlier_year: int | None = None,
) -> float:
    """Refuse a year whose matrix cells do not all move by one ratio.

    The 2009 workbook's note says the thresholds are updated each year by
    the change in the average annual CPI-U, and every cell of the real
    2003-2012 workbooks moves year on year by one ratio to within about
    $1 of whole-dollar rounding; a cell read from the wrong row or column
    moves by far more.  ``earlier`` is the matrix of ``earlier_year``
    (default ``year - 1``); across a gap of several years (the Track M
    capture before 2003) a common ratio still holds, since each year's
    update multiplies every cell by the same factor.  Returns the median
    ratio of ``year`` to ``earlier_year``.
    """

    if earlier_year is None:
        earlier_year = year - 1

    pairs = [
        (earlier[key][k], later[key][k])
        for key in adjusted_poverty.THRESHOLD_ROW_KEYS
        for k in later[key]
    ]
    ratios = sorted(new / old for old, new in pairs)
    middle = len(ratios) // 2
    ratio = (
        ratios[middle]
        if len(ratios) % 2
        else 0.5 * (ratios[middle - 1] + ratios[middle])
    )
    worst = max(abs(new - ratio * old) for old, new in pairs)
    if worst > _CROSS_YEAR_TOLERANCE:
        raise ValueError(
            f"{year}: a matrix cell is ${worst:.2f} from {earlier_year}'s "
            f"times the common ratio {ratio:.6f} (tolerance "
            f"${_CROSS_YEAR_TOLERANCE:.0f})"
        )
    return ratio


def read_workbook_rows(
    path: Path, year: int | None = None
) -> tuple[str, list[list[Any]]]:
    """The table worksheet's title and cell values (formulas as text).

    A workbook holds one worksheet, except that a year listed in
    :data:`EMPTY_EXTRA_SHEETS` may carry exactly the worksheets named there
    after the table's, each without any cell value (whitespace-only cells
    count as blank).  Any other extra worksheet is refused.
    """

    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        sheets = workbook.worksheets
        allowed = EMPTY_EXTRA_SHEETS.get(year, ()) if year is not None else ()
        extra = tuple(sheet.title for sheet in sheets[1:])
        if len(sheets) != 1 and extra != allowed:
            raise ValueError(
                f"{path.name}: {len(sheets)} worksheets {[s.title for s in sheets]}, "
                f"expected 1"
                + (f" plus the empty {list(allowed)}" if allowed else "")
            )
        for sheet in sheets[1:]:
            filled = [
                cell.coordinate
                for row in sheet.iter_rows()
                for cell in row
                if not _blank(cell.value)
            ]
            if filled:
                raise ValueError(
                    f"{path.name}: worksheet {sheet.title!r} is not empty "
                    f"({filled[:3]})"
                )
        sheet = sheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        return sheet.title, rows
    finally:
        workbook.close()


def _parse_years(
    census_dir: Path,
    years: tuple[int, ...],
    expected_sha256: dict[str, str] | None,
) -> dict[str, Any]:
    """Parse ``years``' workbooks in ``census_dir`` in order, pinned first."""

    census_dir = Path(census_dir)
    weighted: dict[str, Any] = {}
    all_ages: dict[str, Any] = {}
    matrix: dict[str, Any] = {}
    sources: dict[str, Any] = {}
    ratios: dict[str, float] = {}
    previous: tuple[int, dict[str, dict[int, int]]] | None = None
    if list(years) != sorted(set(years)):
        raise ValueError(f"years {years} are not strictly increasing")
    for year in years:
        name = f"thresh{year % 100:02d}.xlsx"
        path = census_dir / name
        if not path.is_file():
            raise FileNotFoundError(f"missing Census file {path}")
        digest = _sha256(path)
        if expected_sha256 is not None and digest != expected_sha256.get(name):
            raise ValueError(
                f"{path} sha256 {digest} != pinned "
                f"{expected_sha256.get(name)}: not the captured Census file"
            )
        sheet, rows = read_workbook_rows(path, year)
        parsed = parse_threshold_rows(rows, year)
        if previous is not None:
            earlier_year, earlier = previous
            ratios[f"{earlier_year}-{year}"] = round(
                check_matrix_moves_together(
                    earlier, parsed["matrix"], year, earlier_year
                ),
                6,
            )
        previous = (year, parsed["matrix"])
        weighted[str(year)] = parsed["weighted_average"]
        all_ages[str(year)] = parsed["all_ages"]
        matrix[str(year)] = {
            key: {str(k): v for k, v in cells.items()}
            for key, cells in parsed["matrix"].items()
        }
        source: dict[str, Any] = {
            "file": name,
            "url": CENSUS_URL_BASE + name,
            "sha256": digest,
            "bytes": path.stat().st_size,
            "sheet": sheet,
            "title": parsed["title"],
            "source_line": parsed["source_line"],
        }
        # The CPS ASEC years keep the record Track U's capture pins; the
        # earlier notes name another survey, recorded as printed.
        if parsed["survey"] == "cps_asec":
            source["cps_asec_year"] = parsed["cps_asec_year"]
        else:
            source["note_survey"] = parsed["survey"]
            source["note_survey_year"] = parsed["survey_year"]
        if parsed["revision_line"] is not None:
            source["revision_line"] = parsed["revision_line"]
        if name in ARCHIVE_RETRIEVALS:
            source["retrieval"] = ARCHIVE_RETRIEVALS[name]
        sources[str(year)] = source
    return {
        "weighted": weighted,
        "all_ages": all_ages,
        "matrix": matrix,
        "sources": sources,
        "ratios": ratios,
    }


def build_threshold_capture(
    census_dir: Path,
    *,
    expected_sha256: dict[str, str] | None = CENSUS_WORKBOOK_SHA256,
) -> dict[str, Any]:
    """Parse the nine workbooks in ``census_dir`` into the capture.

    Each workbook's SHA-256 must equal ``expected_sha256`` (the pinned
    Census files) before it is parsed; ``None`` skips the pin and is for
    tests on INVENTED workbooks only (the command line always pins).
    """

    parsed = _parse_years(census_dir, YEARS, expected_sha256)
    weighted = parsed["weighted"]
    all_ages = parsed["all_ages"]
    matrix = parsed["matrix"]
    sources = parsed["sources"]
    ratios = parsed["ratios"]
    return {
        "schema_version": adjusted_poverty.THRESHOLDS_SCHEMA_VERSION,
        "description": (
            "U.S. Census Bureau poverty thresholds for income years "
            "2004-2012, as printed in the Census historical threshold "
            "workbooks: weighted averages by family size (with the "
            "under-65 and 65-and-over rows for one and two persons), the "
            "size-1 and size-2 weighted averages over both ages, and the "
            "size-by-related-children matrix, in dollars"
        ),
        "years": [YEARS[0], YEARS[-1]],
        "sources": sources,
        "generated_by": "scripts/capture_track_u_parameters.py --census-dir",
        "checks": {
            "workbook_sha256": (
                "pinned (CENSUS_WORKBOOK_SHA256)"
                if expected_sha256 is not None
                else "not pinned (INVENTED workbooks)"
            ),
            "layout": (
                "title names the year, source and note name the next "
                "year's CPS ASEC, header and thirteen row labels as "
                "printed, whole-dollar cells in the expected columns, "
                "nothing else in the sheet"
            ),
            "within_year": (
                "65-and-over below under-65 (sizes 1, 2); size-1 and "
                "size-2 all-ages averages between their age rows; "
                "weighted averages rise with size from two; each weighted "
                "average within its row's matrix cells"
            ),
            "cross_year": (
                "every matrix cell moves by one ratio from the previous "
                f"year within ${_CROSS_YEAR_TOLERANCE:.0f}"
            ),
            "matrix_ratio_by_year_pair": ratios,
        },
        "weighted_average": weighted,
        "weighted_average_all_ages": all_ages,
        "matrix": matrix,
    }


def build_track_m_threshold_capture(
    census_dir: Path,
    *,
    expected_sha256: dict[str, str] | None = CENSUS_WORKBOOK_SHA256,
) -> dict[str, Any]:
    """Parse the thirty-five workbooks 1982-2022 into Track M's capture.

    The same parser and checks as :func:`build_threshold_capture`, over
    ``min_benefit_track_m.thresholds.TRACK_M_THRESHOLD_YEARS`` (1982, 1986,
    1988, 1989, 1991, 1992 and 1994-2022; the years before 2003 are those
    M4's structural count shows the in-window records need), with each
    year's layout variant recorded (its worksheets, the unit its weighted
    averages are printed to, the survey its note names, the noun of its
    size rows and any revision line) and the years within 1982-2022 that
    are not captured listed.  The cross-year check runs between
    consecutive captured years, across the gaps too.  Track M reads
    ``weighted_average[year]["one_65_plus"]``: the Census weighted average
    for one person aged 65 and over (the M1 specification's threshold,
    section 7).
    """

    years = track_m_thresholds.TRACK_M_THRESHOLD_YEARS
    parsed = _parse_years(census_dir, years, expected_sha256)
    layouts = {
        str(year): {
            "worksheets": 1 + len(EMPTY_EXTRA_SHEETS.get(year, ())),
            "empty_extra_worksheets": list(EMPTY_EXTRA_SHEETS.get(year, ())),
            "weighted_average_unit_dollars": WEIGHTED_AVERAGE_UNIT.get(
                year, 1
            ),
            "note_survey": note_survey(year),
            "size_row_noun": (
                "persons" if year in PERSONS_ROW_LABEL_YEARS else "people"
            ),
            "revision_line": REVISION_LINES.get(year),
        }
        for year in years
    }
    not_captured = sorted(set(range(years[0], years[-1] + 1)) - set(years))
    return {
        "schema_version": adjusted_poverty.THRESHOLDS_SCHEMA_VERSION,
        "description": (
            "U.S. Census Bureau poverty thresholds for income years "
            f"{years[0]}-{years[-1]} (the captured_years only), as printed "
            "in the Census historical threshold workbooks, for Track M "
            "(DynaSim exercise 4, the minimum benefit; Python rules, not "
            "Axiom), which reads the weighted average for one person aged "
            "65 and over (one_65_plus). Same content and schema as the "
            "Track U capture: weighted averages by family size (with the "
            "under-65 and 65-and-over rows for one and two persons), the "
            "size-1 and size-2 weighted averages over both ages, and the "
            "size-by-related-children matrix, in dollars"
        ),
        "years": [years[0], years[-1]],
        "captured_years": list(years),
        "years_not_captured": not_captured,
        "decision_records": {
            "d194": "thresh03-thresh12 staged 2026-09-24 (exercise 2)",
            "d279": (
                "thresh13-thresh22 downloaded 2026-09-25 for Track M; "
                "capture, hash and pin like the 2003-2012 capture"
            ),
            "d279_earlier_years": (
                "thresh82, thresh86, thresh88, thresh89, thresh91, thresh92 "
                "and thresh94-thresh02 staged 2026-09-25 under d279 (any "
                "earlier year the build proves it needs: M4's structural "
                "count of the threshold years the in-window records need); "
                "thresh95 from the Internet Archive's copy of its Census URL"
            ),
        },
        "sources": parsed["sources"],
        "generated_by": (
            "scripts/capture_track_u_parameters.py --track-m-census-dir"
        ),
        "checks": {
            "workbook_sha256": (
                "pinned (CENSUS_WORKBOOK_SHA256)"
                if expected_sha256 is not None
                else "not pinned (INVENTED workbooks)"
            ),
            "layout": (
                "title names the year, source and note name the next "
                "year's survey (the March CPS 1982-2000, the CPS ADS 2001, "
                "the CPS ASEC from 2002), header and thirteen row labels "
                "as printed (2001's size rows say persons), whole-dollar "
                "cells in the expected columns, nothing else in the "
                "table's sheet but 1982's and 2000's revision lines below "
                "the note; 2019 alone may carry the empty worksheets "
                "Sheet2 and Sheet3"
            ),
            "within_year": (
                "65-and-over below under-65 (sizes 1, 2); size-1 and "
                "size-2 all-ages averages between their age rows; "
                "weighted averages rise with size from two; each weighted "
                "average within its row's matrix cells, or, in 2022, whose "
                "workbook prints weighted averages rounded to $10, a "
                "multiple of $10 within $5 of them"
            ),
            "cross_year": (
                "every matrix cell moves by one ratio from the previous "
                f"captured year within ${_CROSS_YEAR_TOLERANCE:.0f}"
            ),
            "matrix_ratio_by_year_pair": parsed["ratios"],
            "layout_by_year": layouts,
        },
        "weighted_average": parsed["weighted"],
        "weighted_average_all_ages": parsed["all_ages"],
        "matrix": parsed["matrix"],
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--ssi", action="store_true")
    parser.add_argument("--pe-us-dir", type=Path, default=None)
    parser.add_argument("--census-dir", type=Path, default=None)
    parser.add_argument(
        "--ssi-output", type=Path, default=adjusted_poverty.SSI_PARAMETERS_PATH
    )
    parser.add_argument(
        "--census-output", type=Path, default=adjusted_poverty.THRESHOLDS_PATH
    )
    parser.add_argument("--track-m-census-dir", type=Path, default=None)
    parser.add_argument(
        "--track-m-output",
        type=Path,
        default=track_m_thresholds.TRACK_M_THRESHOLDS_PATH,
    )
    args = parser.parse_args(argv)
    if (
        not args.ssi
        and args.census_dir is None
        and args.track_m_census_dir is None
    ):
        parser.error("pass --ssi, --census-dir and/or --track-m-census-dir")
    if args.ssi:
        capture = build_ssi_capture(args.pe_us_dir)
        args.ssi_output.write_text(
            json.dumps(capture, indent=2) + "\n", encoding="utf-8"
        )
        print(args.ssi_output, _sha256(args.ssi_output))
    if args.census_dir is not None:
        capture = build_threshold_capture(args.census_dir)
        args.census_output.write_text(
            json.dumps(capture, indent=2) + "\n", encoding="utf-8"
        )
        print(args.census_output, _sha256(args.census_output))
    if args.track_m_census_dir is not None:
        capture = build_track_m_threshold_capture(args.track_m_census_dir)
        args.track_m_output.write_text(
            json.dumps(capture, indent=2) + "\n", encoding="utf-8"
        )
        print(args.track_m_output, _sha256(args.track_m_output))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
